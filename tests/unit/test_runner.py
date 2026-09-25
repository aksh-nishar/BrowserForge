import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from browserforge.engine.driver import DriverTimeoutError
from browserforge.engine.recipe import STEP_TYPES, Recipe, Step
from browserforge.engine.runner import HANDLERS, RunLog, run_recipe, total_steps
from fixtures.fake_driver import FakePageDriver

URL = "https://portal.example.com"
FIXED_NOW = datetime(2026, 9, 25, 10, 15, 30)
PASSWORD = "S3cret!pw"


def step(step_type: str, name: str | None = None, **config: Any) -> Step:
    return Step(type=step_type, name=name, config=config)


def recipe(*steps: Step, **kwargs: Any) -> Recipe:
    return Recipe(url=URL, steps=list(steps), **kwargs)


def run(
    rec: Recipe, fake: FakePageDriver | None, job_dir: Path, **kwargs: Any
) -> tuple[RunLog, FakePageDriver]:
    driver = fake or FakePageDriver()
    kwargs.setdefault("now", lambda: FIXED_NOW)
    return run_recipe(rec, driver, job_dir, **kwargs), driver


@pytest.fixture
def job_dir(tmp_path: Path) -> Path:
    return tmp_path / "job-1"  # not created: the runner must create it


def read_run_log(job_dir: Path) -> dict[str, Any]:
    data: dict[str, Any] = json.loads((job_dir / "run_log.json").read_text(encoding="utf-8"))
    return data


# ---- happy path ----


def test_mixed_steps_all_run_in_order_and_succeed(job_dir: Path) -> None:
    rec = recipe(
        step("fill", selector="#user", value="bob"),
        step("click", selector="#go"),
        step("screenshot", name="done"),
    )
    log, fake = run(rec, None, job_dir)
    assert [s.type for s in log.steps] == ["goto", "fill", "click", "screenshot"]
    assert [s.index for s in log.steps] == [1, 2, 3, 4]
    assert all(s.success for s in log.steps)
    assert log.success is True
    assert log.stopped_early is False
    assert [name for name, _ in fake.calls][:3] == ["goto", "fill", "click"]


def test_start_page_is_opened_first(job_dir: Path) -> None:
    _, fake = run(recipe(step("click", selector="#go")), None, job_dir)
    assert fake.calls[0] == ("goto", (URL,))


def test_empty_recipe_only_opens_the_start_page(job_dir: Path) -> None:
    log, fake = run(recipe(), None, job_dir)
    assert [s.type for s in log.steps] == ["goto"]
    assert log.success is True
    assert fake.calls == [("goto", (URL,))]


def test_login_block_runs_after_goto_and_before_steps(job_dir: Path) -> None:
    login = {
        "username_selector": "#u",
        "password_selector": "#p",
        "submit_selector": "#s",
        "username": "tom",
        "password": PASSWORD,
        "success_selector": "#dash",
    }
    log, fake = run(recipe(step("click", selector="#report"), login=login), None, job_dir)
    assert [s.type for s in log.steps] == ["goto", "login", "click"]
    assert log.success is True
    assert [name for name, _ in fake.calls] == [
        "goto",
        "fill",
        "fill",
        "click",
        "wait_for_selector",
        "click",
    ]


def test_credentials_never_appear_in_run_log_file(job_dir: Path) -> None:
    login = {
        "username_selector": "#u",
        "password_selector": "#p",
        "submit_selector": "#s",
        "username": "tom",
        "password": PASSWORD,
        "success_selector": "#dash",
    }
    fake = FakePageDriver()
    fake.set_missing("#dash")  # wrong password -> login failure recorded in the log
    log, _ = run(recipe(login=login), fake, job_dir)
    assert not log.success
    text = (job_dir / "run_log.json").read_text(encoding="utf-8")
    assert PASSWORD not in text
    assert "Login failed" in text


# ---- run_log.json ----


def test_run_log_json_is_written_with_expected_structure(job_dir: Path) -> None:
    log, _ = run(recipe(step("click", selector="#go")), None, job_dir, recipe_name="demo")
    data = read_run_log(job_dir)
    assert data["job_id"] == "job-1"
    assert data["recipe_name"] == "demo"
    assert data["url"] == URL
    assert data["success"] is True
    assert data["started_at"] == "2026-09-25T10:15:30"
    assert data["finished_at"] == "2026-09-25T10:15:30"
    assert data["stopped_early"] is False
    assert data["skipped_steps"] == 0
    assert data["output_file"] is None
    assert [s["type"] for s in data["steps"]] == ["goto", "click"]
    first = data["steps"][0]
    assert set(first) == {
        "index",
        "type",
        "name",
        "success",
        "error",
        "duration_ms",
        "data",
        "attempts",
        "retry_errors",
        "error_kind",
        "human_error",
        "failure_screenshot",
    }
    assert first["name"] == "open_start_page"
    assert isinstance(first["duration_ms"], int)
    assert first["duration_ms"] >= 0
    assert data["duration_ms"] >= 0
    assert data["steps"][1]["success"] is True
    assert log.job_id == "job-1"


def test_job_dir_is_created_and_str_paths_are_accepted(tmp_path: Path) -> None:
    target = tmp_path / "deep" / "er" / "job-x"
    run_recipe(recipe(), FakePageDriver(), str(target), now=lambda: FIXED_NOW)
    assert (target / "run_log.json").is_file()


def test_job_id_defaults_to_folder_name_and_can_be_overridden(tmp_path: Path) -> None:
    default_log, _ = run(recipe(), None, tmp_path / "abc123")
    custom_log, _ = run(recipe(), None, tmp_path / "other", job_id="custom-id")
    assert default_log.job_id == "abc123"
    assert custom_log.job_id == "custom-id"


def test_non_ascii_content_survives_in_the_log(job_dir: Path) -> None:
    run(recipe(step("click", name="बटन", selector="#go")), None, job_dir)
    assert "बटन" in (job_dir / "run_log.json").read_text(encoding="utf-8")


# ---- failure handling ----


def test_failure_stops_the_run_by_default(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_missing("#missing")
    rec = recipe(
        step("click", selector="#a"),
        step("click", selector="#missing"),
        step("click", selector="#never"),
        step("click", selector="#never2"),
    )
    log, _ = run(rec, fake, job_dir)
    assert [s.success for s in log.steps] == [True, True, False]
    assert log.stopped_early is True
    assert log.skipped_steps == 2
    assert log.success is False
    assert log.steps[-1].error is not None
    assert "#missing" in log.steps[-1].error
    assert fake.clicks == ["#a"]
    assert read_run_log(job_dir)["skipped_steps"] == 2


def test_continue_on_error_runs_remaining_steps(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_missing("#missing")
    rec = recipe(
        step("click", selector="#missing"),
        step("click", selector="#after"),
        continue_on_error=True,
    )
    log, _ = run(rec, fake, job_dir)
    assert [s.success for s in log.steps] == [True, False, True]
    assert log.stopped_early is False
    assert log.skipped_steps == 0
    assert log.success is False  # any failed step makes the run unsuccessful
    assert fake.clicks == ["#after"]


def test_failed_start_page_stops_before_any_step(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.fail_on("goto", DriverTimeoutError("net::ERR_TIMED_OUT"))
    log, _ = run(recipe(step("click", selector="#a")), fake, job_dir)
    assert [s.type for s in log.steps] == ["goto"]
    assert log.steps[0].success is False
    assert log.skipped_steps == 1


def test_unknown_step_type_is_a_failed_step_not_a_crash(job_dir: Path) -> None:
    rec = recipe(
        step("click", selector="#a"),
        Step(type="teleport", config={}),
        step("click", selector="#b"),
        continue_on_error=True,
    )
    log, fake = run(rec, None, job_dir)
    assert [s.success for s in log.steps] == [True, True, False, True]
    assert log.steps[2].error is not None
    assert "teleport" in log.steps[2].error
    assert fake.clicks == ["#a", "#b"]


def test_unknown_step_type_stops_run_without_continue_flag(job_dir: Path) -> None:
    rec = recipe(Step(type="teleport"), step("click", selector="#b"))
    log, fake = run(rec, None, job_dir)
    assert log.stopped_early is True
    assert fake.clicks == []
    assert (job_dir / "run_log.json").is_file()


def test_unexpected_exception_in_an_action_is_contained(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.fail_on("click", RuntimeError("driver bug"))
    rec = recipe(step("click", selector="#a"), step("check", selector="#b"), continue_on_error=True)
    log, _ = run(rec, fake, job_dir)
    assert [s.success for s in log.steps] == [True, False, True]
    assert log.steps[1].error is not None
    assert "RuntimeError" in log.steps[1].error
    assert "driver bug" in log.steps[1].error
    assert (job_dir / "run_log.json").is_file()


def test_bad_step_config_is_reported_via_the_step(job_dir: Path) -> None:
    log, fake = run(recipe(step("fill", selector="#a")), None, job_dir)  # value missing
    assert log.steps[-1].success is False
    assert log.steps[-1].error is not None
    assert "value" in log.steps[-1].error
    assert fake.fills == {}


# ---- step data, screenshots, downloads ----


def test_screenshot_steps_get_index_based_files_inside_the_job(job_dir: Path) -> None:
    rec = recipe(step("screenshot", name="one"), step("screenshot", name="two"))
    log, _ = run(rec, None, job_dir)
    paths = [Path(s.data["path"]) for s in log.steps if s.type == "screenshot"]
    assert len(set(paths)) == 2
    assert all(p.parent == job_dir / "screenshots" and p.exists() for p in paths)
    assert paths[0].name.startswith("02_one_")
    assert paths[1].name.startswith("03_two_")


def test_download_step_saves_into_job_downloads(tmp_path: Path, job_dir: Path) -> None:
    source = tmp_path / "tmp" / "report.csv"
    source.parent.mkdir()
    source.write_text("a,b")
    fake = FakePageDriver()
    fake.queue_download(str(source))
    log, _ = run(recipe(step("download", selector="#export")), fake, job_dir)
    assert log.success
    assert (job_dir / "downloads" / "report.csv").read_text() == "a,b"


# ---- tables -> output.xlsx ----


def test_extracted_table_is_exported_to_output_xlsx(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_table("#tbl", [["name", "age"], ["Ann", "30"], ["Bob", "25"]])
    log, _ = run(recipe(step("extract_table", name="People", selector="#tbl")), fake, job_dir)
    assert log.output_file == str(job_dir / "output.xlsx")
    ws = load_workbook(job_dir / "output.xlsx")["People"]
    assert [[c.value for c in row] for row in ws.iter_rows()] == [
        ["name", "age"],
        ["Ann", "30"],
        ["Bob", "25"],
    ]


def test_table_rows_are_summarised_not_embedded_in_run_log(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_table("#tbl", [["name"], ["Ann"], ["Bob"], ["Cy"]])
    run(recipe(step("extract_table", name="People", selector="#tbl")), fake, job_dir)
    data = read_run_log(job_dir)
    table_step = data["steps"][1]
    assert table_step["data"] == {"row_count": 3}
    assert "Ann" not in (job_dir / "run_log.json").read_text(encoding="utf-8")
    assert data["output_file"] == str(job_dir / "output.xlsx")


def test_paginate_step_exports_all_pages_and_reports_truncation(job_dir: Path) -> None:
    fake = FakePageDriver()
    header = ["id"]
    fake.set_pages("#tbl", "#next", [[header, ["1"]], [header, ["2"]], [header, ["3"]]])
    rec = recipe(step("paginate", name="All", selector="#tbl", next_selector="#next", max_pages=2))
    log, _ = run(rec, fake, job_dir)
    paginate_log = log.steps[1]
    assert paginate_log.data == {"row_count": 2, "pages": 2, "truncated": True}
    ws = load_workbook(job_dir / "output.xlsx")["All"]
    assert [row[0].value for row in ws.iter_rows()] == ["id", "1", "2"]


def test_two_tables_become_two_sheets_even_with_the_same_step_name(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_table("#a", [["x"], ["1"]])
    fake.set_table("#b", [["y"], ["2"]])
    rec = recipe(
        step("extract_table", name="Data", selector="#a"),
        step("extract_table", name="data", selector="#b"),
    )
    run(rec, fake, job_dir)
    assert load_workbook(job_dir / "output.xlsx").sheetnames == ["Data", "data_2"]


def test_unnamed_table_step_gets_a_sheet_named_from_type_and_position(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_table("#a", [["x"], ["1"]])
    run(recipe(step("extract_table", selector="#a")), fake, job_dir)
    assert load_workbook(job_dir / "output.xlsx").sheetnames == ["extract_table_2"]


def test_no_table_steps_means_no_excel_file(job_dir: Path) -> None:
    log, _ = run(recipe(step("click", selector="#a")), None, job_dir)
    assert log.output_file is None
    assert not (job_dir / "output.xlsx").exists()


def test_failed_table_step_produces_no_sheet(job_dir: Path) -> None:
    log, _ = run(recipe(step("extract_table", selector="#nope")), None, job_dir)  # not scripted
    assert not log.success
    assert not (job_dir / "output.xlsx").exists()


def test_excel_write_failure_is_recorded_and_run_log_still_written(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_table("#a", [["x"], ["1"]])
    (job_dir / "output.xlsx").mkdir(parents=True)  # a directory where the file must go
    log, _ = run(recipe(step("extract_table", name="T", selector="#a")), fake, job_dir)
    assert log.success is False
    assert log.steps[-1].type == "export_excel"
    assert log.steps[-1].error is not None
    assert "output.xlsx" in log.steps[-1].error
    assert log.output_file is None
    assert read_run_log(job_dir)["success"] is False


# ---- progress callbacks (used by the job store for live status) ----


def test_progress_callbacks_fire_in_order_for_every_step(job_dir: Path) -> None:
    events: list[tuple[str, int, str]] = []
    rec = recipe(step("click", name="a", selector="#a"), step("check", name="b", selector="#b"))
    run(
        rec,
        None,
        job_dir,
        on_step_start=lambda i, total, s: events.append(
            ("start", i, f"{total}:{s.name or s.type}")
        ),
        on_step_end=lambda log: events.append(("end", log.index, log.name or log.type)),
    )
    assert events == [
        ("start", 1, "3:open_start_page"),
        ("end", 1, "open_start_page"),
        ("start", 2, "3:a"),
        ("end", 2, "a"),
        ("start", 3, "3:b"),
        ("end", 3, "b"),
    ]


def test_end_callback_receives_the_failure_and_stops_follow_ups(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_missing("#a")
    ended: list[bool] = []
    run(
        recipe(step("click", selector="#a"), step("click", selector="#b")),
        fake,
        job_dir,
        on_step_end=lambda log: ended.append(log.success),
    )
    assert ended == [True, False]  # nothing reported for the skipped step


def test_export_failure_is_reported_to_the_end_callback(job_dir: Path) -> None:
    fake = FakePageDriver()
    fake.set_table("#a", [["x"], ["1"]])
    (job_dir / "output.xlsx").mkdir(parents=True)
    ended: list[str] = []
    run(
        recipe(step("extract_table", selector="#a")),
        fake,
        job_dir,
        on_step_end=lambda log: ended.append(log.type),
    )
    assert ended[-1] == "export_excel"


def test_a_crashing_progress_callback_never_breaks_the_run(job_dir: Path) -> None:
    def explode(*_args: object) -> None:
        raise RuntimeError("callback bug")

    log, fake = run(
        recipe(step("click", selector="#a")),
        None,
        job_dir,
        on_step_start=explode,
        on_step_end=explode,
    )
    assert log.success is True
    assert fake.clicks == ["#a"]


# ---- wiring ----


LOGIN_BLOCK = {
    "username_selector": "#u",
    "password_selector": "#p",
    "submit_selector": "#s",
    "username": "tom",
    "password": "pw",
    "success_selector": "#dash",
}


@pytest.mark.parametrize(
    ("steps", "login", "expected"),
    [
        (0, False, 1),  # just the start page
        (3, False, 4),
        (0, True, 2),  # start page + login
        (3, True, 5),
    ],
)
def test_total_steps_matches_the_rows_a_real_run_produces(
    job_dir: Path, steps: int, login: bool, expected: int
) -> None:
    rec = recipe(
        *[step("click", selector=f"#b{i}") for i in range(steps)],
        login=LOGIN_BLOCK if login else None,
    )
    log, _ = run(rec, None, job_dir)
    assert total_steps(rec) == expected
    assert len(log.steps) == expected  # what the dashboard's "(N steps)" promises


def test_every_recipe_step_type_has_a_handler() -> None:
    assert set(HANDLERS) == set(STEP_TYPES)
