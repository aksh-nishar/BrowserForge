from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from browserforge.engine.excel_export import (
    MAX_CELL_CHARS,
    MAX_COLUMN_WIDTH,
    MIN_COLUMN_WIDTH,
    build_workbook,
    column_width,
    safe_sheet_name,
    save_workbook,
)

ROWS = [
    {"name": "Ann", "age": "30"},
    {"name": "Bob", "age": "25"},
    {"name": "Cy", "age": "41"},
]


def sheet_values(ws: Any) -> list[list[Any]]:
    return [[cell.value for cell in row] for row in ws.iter_rows()]


# ---- structure ----


def test_single_sheet_three_rows() -> None:
    wb = build_workbook({"People": ROWS})
    assert wb.sheetnames == ["People"]
    assert sheet_values(wb["People"]) == [
        ["name", "age"],
        ["Ann", "30"],
        ["Bob", "25"],
        ["Cy", "41"],
    ]


def test_header_row_is_bold_and_data_rows_are_not() -> None:
    ws = build_workbook({"People": ROWS})["People"]
    assert all(cell.font.bold for cell in ws[1])
    assert not any(cell.font.bold for cell in ws[2])


def test_header_row_is_frozen() -> None:
    assert build_workbook({"People": ROWS})["People"].freeze_panes == "A2"


def test_multiple_sheets_keep_names_and_order() -> None:
    wb = build_workbook({"First": ROWS, "Second": [{"x": "1"}]})
    assert wb.sheetnames == ["First", "Second"]
    assert sheet_values(wb["Second"]) == [["x"], ["1"]]


def test_no_default_blank_sheet_is_left_behind() -> None:
    assert "Sheet" not in build_workbook({"People": ROWS}).sheetnames


# ---- empty input ----


def test_empty_rows_creates_sheet_with_no_data() -> None:
    wb = build_workbook({"Empty": []})
    assert wb.sheetnames == ["Empty"]
    assert sheet_values(wb["Empty"]) == []
    assert wb["Empty"]["A1"].value is None


def test_empty_mapping_yields_one_empty_sheet_that_can_be_saved(tmp_path: Path) -> None:
    wb = build_workbook({})
    assert wb.sheetnames == ["Sheet"]
    save_workbook(wb, tmp_path / "empty.xlsx")  # would raise if the workbook had no sheets


def test_rows_that_are_all_empty_dicts_write_nothing() -> None:
    ws = build_workbook({"S": [{}, {}]})["S"]
    assert ws.max_row == 1
    assert ws["A1"].value is None


# ---- columns ----


def test_rows_with_different_keys_get_union_of_columns_in_first_seen_order() -> None:
    wb = build_workbook({"S": [{"a": "1", "b": "2"}, {"b": "3", "c": "4"}]})
    assert sheet_values(wb["S"]) == [["a", "b", "c"], ["1", "2", None], [None, "3", "4"]]


def test_none_values_become_blank_cells() -> None:
    ws = build_workbook({"S": [{"a": None, "b": "x"}]})["S"]
    assert ws["A2"].value is None
    assert ws["B2"].value == "x"


# ---- autofit ----


def test_column_width_reflects_longest_value() -> None:
    assert (
        column_width(["short", "a considerably longer value"])
        == len("a considerably longer value") + 2
    )


def test_column_width_has_a_minimum() -> None:
    assert column_width(["a"]) == MIN_COLUMN_WIDTH
    assert column_width([]) == MIN_COLUMN_WIDTH
    assert column_width([None]) == MIN_COLUMN_WIDTH


def test_column_width_is_capped() -> None:
    assert column_width(["x" * 500]) == MAX_COLUMN_WIDTH


def test_column_width_uses_longest_line_of_multiline_text() -> None:
    assert column_width(["ab\n" + "c" * 20 + "\nde"]) == 22


def test_column_width_handles_numbers() -> None:
    assert column_width([1234567890123]) == 15


def test_widths_are_applied_to_sheet_columns() -> None:
    ws = build_workbook({"S": [{"short": "x", "long": "y" * 40}]})["S"]
    assert ws.column_dimensions["A"].width == MIN_COLUMN_WIDTH
    assert ws.column_dimensions["B"].width == 42


def test_header_text_counts_toward_width() -> None:
    ws = build_workbook({"S": [{"a_rather_long_header_name": "1"}]})["S"]
    assert ws.column_dimensions["A"].width == len("a_rather_long_header_name") + 2


# ---- special characters and hostile values ----


def test_unicode_and_newlines_are_written_correctly() -> None:
    ws = build_workbook({"S": [{"नाम": "रवि", "note": "line1\nline2", "e": "café ☕"}]})["S"]
    assert sheet_values(ws) == [["नाम", "note", "e"], ["रवि", "line1\nline2", "café ☕"]]


def test_formula_like_strings_stay_text() -> None:
    payload = '=HYPERLINK("http://evil.example","click")'
    ws = build_workbook({"S": [{"a": payload}]})["S"]
    assert ws["A2"].value == payload
    assert ws["A2"].data_type == "s"


def test_formula_like_strings_stay_text_after_save_and_reload(tmp_path: Path) -> None:
    path = save_workbook(build_workbook({"S": [{"a": "=1+1"}]}), tmp_path / "f.xlsx")
    cell = load_workbook(path)["S"]["A2"]
    assert cell.value == "=1+1"
    assert cell.data_type != "f"


def test_control_characters_are_stripped_instead_of_crashing() -> None:
    ws = build_workbook({"S": [{"a": "bad\x00\x07value", "b\x0b": "ok"}]})["S"]
    assert ws["A2"].value == "badvalue"
    assert ws["B1"].value == "b"


def test_overlong_string_is_truncated_to_the_excel_limit() -> None:
    ws = build_workbook({"S": [{"a": "x" * (MAX_CELL_CHARS + 500)}]})["S"]
    assert len(ws["A2"].value) == MAX_CELL_CHARS


def test_very_long_string_within_limit_is_written_in_full() -> None:
    text = "y" * 20_000
    assert build_workbook({"S": [{"a": text}]})["S"]["A2"].value == text


def test_numbers_and_booleans_keep_their_type() -> None:
    ws = build_workbook({"S": [{"i": 5, "f": 2.5, "b": True}]})["S"]
    assert (ws["A2"].value, ws["B2"].value, ws["C2"].value) == (5, 2.5, True)


def test_other_objects_are_stringified() -> None:
    ws = build_workbook({"S": [{"a": ["x", "y"], "b": {"k": 1}}]})["S"]
    assert ws["A2"].value == "['x', 'y']"
    assert ws["B2"].value == "{'k': 1}"


# ---- sheet names ----


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Report", "Report"),
        ("a/b\\c:d*e?f[g]h", "a_b_c_d_e_f_g_h"),
        ("", "Sheet"),
        ("   ", "Sheet"),
        ("'quoted'", "quoted"),
        ("x" * 50, "x" * 31),
    ],
)
def test_safe_sheet_name_rules(raw: str, expected: str) -> None:
    assert safe_sheet_name(raw) == expected


def test_safe_sheet_name_dedupes_case_insensitively_within_31_chars() -> None:
    assert safe_sheet_name("Data", taken=["data"]) == "Data_2"
    assert safe_sheet_name("Data", taken=["data", "Data_2"]) == "Data_3"
    long = safe_sheet_name("x" * 40, taken=["x" * 31])
    assert len(long) == 31
    assert long.endswith("_2")


def test_colliding_sheet_names_do_not_lose_data() -> None:
    wb = build_workbook({"A/B": [{"v": "1"}], "A_B": [{"v": "2"}]})
    assert wb.sheetnames == ["A_B", "A_B_2"]
    assert wb["A_B_2"]["A2"].value == "2"


def test_invalid_sheet_name_does_not_crash_build() -> None:
    assert build_workbook({"Q1: [draft]?": ROWS}).sheetnames == ["Q1_ _draft__"]


# ---- disk wrapper ----


def test_save_workbook_writes_a_readable_file_and_creates_folders(tmp_path: Path) -> None:
    target = tmp_path / "jobs" / "job1" / "output.xlsx"
    returned = save_workbook(build_workbook({"People": ROWS}), target)
    assert returned == target
    assert sheet_values(load_workbook(target)["People"])[1] == ["Ann", "30"]


def test_build_workbook_does_not_touch_the_filesystem(tmp_path: Path, monkeypatch: Any) -> None:
    monkeypatch.chdir(tmp_path)
    build_workbook({"People": ROWS})
    assert list(tmp_path.iterdir()) == []
