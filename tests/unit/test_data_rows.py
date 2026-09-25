"""Bulk data: reading sheets, placeholders, mapping, validation and secrets. All pure."""

import io
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from browserforge.engine.data_rows import (
    MAX_ROWS,
    DataError,
    apply_mapping,
    expand_recipe,
    find_env_vars,
    find_placeholders,
    is_secret_name,
    load_rows,
    normalize_key,
    parse_rows,
    redact,
    secret_columns,
    secret_values,
    suggest_mapping,
    validate_env,
    validate_rows,
)
from browserforge.engine.recipe import Recipe, Step, parse_recipe_text

RECIPE_TEXT = """
url: https://x.io/{{site}}
login:
  username: "{{user}}"
  password: "{{env:PORTAL_PASSWORD}}"
steps:
  - type: fill
    selector: "#name"
    value: "{{first_name}}"
  - type: fill
    selector: "#age"
    value: 42
  - type: upload
    selector: "#cv"
    files:
      - "{{cv_path}}"
      - plain.txt
"""


def recipe() -> Recipe:
    return parse_recipe_text(RECIPE_TEXT)


def xlsx_bytes(table: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in table:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---- reading CSV ----


def test_csv_rows_use_the_header_as_keys() -> None:
    rows = parse_rows(b"name,city\nAnn,Pune\nBob,Delhi\n", "people.csv")
    assert rows == [{"name": "Ann", "city": "Pune"}, {"name": "Bob", "city": "Delhi"}]


def test_csv_blank_lines_and_blank_header_columns_are_dropped() -> None:
    rows = parse_rows(b"name,,city\nAnn,x,Pune\n,,\n\nBob,y,Delhi\n", "p.csv")
    assert rows == [{"name": "Ann", "city": "Pune"}, {"name": "Bob", "city": "Delhi"}]


def test_csv_values_are_stripped_and_short_rows_padded() -> None:
    assert parse_rows(b"a,b\n  x  \n", "p.csv") == [{"a": "x", "b": ""}]


def test_csv_with_a_byte_order_mark_and_unicode_is_read() -> None:
    raw = "﻿name\nZoë – 日本\n".encode()
    assert parse_rows(raw, "p.csv") == [{"name": "Zoë – 日本"}]


def test_csv_header_only_gives_zero_rows() -> None:
    assert parse_rows(b"name,city\n", "p.csv") == []


def test_completely_empty_file_is_a_clear_error() -> None:
    with pytest.raises(DataError, match="empty"):
        parse_rows(b"", "p.csv")


def test_duplicate_column_names_are_rejected_naming_them() -> None:
    with pytest.raises(DataError, match="twice: name"):
        parse_rows(b"name,name\nA,B\n", "p.csv")


def test_csv_that_is_not_utf8_explains_how_to_fix_it() -> None:
    with pytest.raises(DataError, match="UTF-8"):
        parse_rows("name\nZoë".encode("latin-1"), "p.csv")


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(DataError, match=r"\.csv or \.xlsx"):
        parse_rows(b"a\n1\n", "people.txt")


def test_too_many_rows_are_rejected() -> None:
    raw = ("n\n" + "x\n" * (MAX_ROWS + 1)).encode()
    with pytest.raises(DataError, match=str(MAX_ROWS)):
        parse_rows(raw, "big.csv")


def test_exactly_the_row_limit_is_accepted() -> None:
    raw = ("n\n" + "x\n" * MAX_ROWS).encode()
    assert len(parse_rows(raw, "big.csv")) == MAX_ROWS


# ---- reading XLSX ----


def test_xlsx_rows_are_read_from_the_first_sheet() -> None:
    raw = xlsx_bytes([["name", "city"], ["Ann", "Pune"], ["Bob", "Delhi"]])
    assert parse_rows(raw, "p.xlsx") == [
        {"name": "Ann", "city": "Pune"},
        {"name": "Bob", "city": "Delhi"},
    ]


def test_xlsx_numbers_and_dates_become_clean_text() -> None:
    raw = xlsx_bytes(
        [
            ["phone", "score", "born", "when", "ok"],
            [9876543210, 4.5, date(2000, 1, 31), datetime(2026, 1, 2, 3, 4, 5), True],
        ]
    )
    (row,) = parse_rows(raw, "p.xlsx")
    assert row["phone"] == "9876543210"  # no ".0" tail
    assert row["score"] == "4.5"
    assert row["born"] == "2000-01-31"
    assert row["when"] == "2026-01-02T03:04:05"
    assert row["ok"] == "true"


def test_xlsx_blank_rows_are_skipped_and_empty_cells_are_blank_text() -> None:
    raw = xlsx_bytes([["a", "b"], [None, None], ["x", None]])
    assert parse_rows(raw, "p.xlsx") == [{"a": "x", "b": ""}]


def test_a_corrupt_xlsx_is_a_plain_error() -> None:
    with pytest.raises(DataError, match="could not be read"):
        parse_rows(b"not really a zip file", "p.xlsx")


def test_load_rows_reads_a_file_and_reports_a_missing_one(tmp_path: Path) -> None:
    file = tmp_path / "p.csv"
    file.write_text("a\n1\n", encoding="utf-8")
    assert load_rows(file) == [{"a": "1"}]
    with pytest.raises(DataError, match="not found"):
        load_rows(tmp_path / "missing.csv")


# ---- placeholders ----


def test_find_placeholders_looks_in_url_login_and_step_settings() -> None:
    assert find_placeholders(recipe()) == {"site", "user", "first_name", "cv_path"}


def test_find_env_vars_is_separate_from_data_placeholders() -> None:
    assert find_env_vars(recipe()) == {"PORTAL_PASSWORD"}


def test_a_recipe_without_placeholders_has_none() -> None:
    plain = parse_recipe_text("url: https://x.io\nsteps:\n  - type: click\n    selector: '#a'\n")
    assert find_placeholders(plain) == set()
    assert find_env_vars(plain) == set()


def test_placeholders_tolerate_inner_spaces_and_ignore_other_braces() -> None:
    text = (
        "url: https://x.io\n"
        "steps:\n"
        "  - type: fill\n"
        '    selector: "#a"\n'
        '    value: "{{ a }} {b} {{}} {{1x}}"\n'
    )
    assert find_placeholders(parse_recipe_text(text)) == {"a"}


def test_expand_recipe_fills_every_field_and_leaves_the_original_alone() -> None:
    original = recipe()
    row = {"site": "home", "user": "ann", "first_name": "Ann", "cv_path": "cv.pdf"}
    expanded = expand_recipe(original, row, {"PORTAL_PASSWORD": "hunter22"})

    assert expanded.url == "https://x.io/home"
    assert expanded.login == {"username": "ann", "password": "hunter22"}
    assert expanded.steps[0].config["value"] == "Ann"
    assert expanded.steps[1].config["value"] == 42  # non-strings are untouched
    assert expanded.steps[2].config["files"] == ["cv.pdf", "plain.txt"]
    assert original.url == "https://x.io/{{site}}"  # the original is unchanged
    assert original.steps[0].config["value"] == "{{first_name}}"


def test_expand_keeps_step_names_types_and_retry_settings() -> None:
    original = Recipe(
        url="https://x.io",
        steps=[Step(type="click", name="go", config={"selector": "#a"}, retries=2)],
        retries=3,
        continue_on_error=True,
    )
    expanded = expand_recipe(original, {})
    assert expanded.steps[0].name == "go"
    assert expanded.steps[0].retries == 2
    assert expanded.retries == 3
    assert expanded.continue_on_error is True


def test_a_missing_value_names_the_placeholder() -> None:
    with pytest.raises(DataError, match="'first_name'"):
        expand_recipe(
            recipe(), {"site": "a", "user": "b", "cv_path": "c"}, {"PORTAL_PASSWORD": "x"}
        )


def test_a_missing_environment_variable_names_it() -> None:
    row = {"site": "a", "user": "b", "first_name": "c", "cv_path": "d"}
    with pytest.raises(DataError, match="PORTAL_PASSWORD"):
        expand_recipe(recipe(), row, {})


def test_substituted_text_is_not_expanded_again() -> None:
    plain = parse_recipe_text(
        'url: https://x.io\nsteps:\n  - type: fill\n    selector: "#a"\n    value: "{{a}}"\n'
    )
    expanded = expand_recipe(plain, {"a": "{{b}}", "b": "boom"})
    assert expanded.steps[0].config["value"] == "{{b}}"


# ---- validation ----


def test_validate_rows_accepts_complete_data() -> None:
    assert validate_rows({"a", "b"}, [{"a": "1", "b": "2"}]) == []


def test_validate_rows_reports_no_rows() -> None:
    assert validate_rows({"a"}, []) == ["The data file has no rows below the header."]


def test_validate_rows_reports_a_missing_column_once() -> None:
    problems = validate_rows({"a", "zip"}, [{"a": "1"}, {"a": "2"}])
    assert problems == ["The data has no column for 'zip'."]


def test_validate_rows_reports_blank_cells_with_row_numbers() -> None:
    problems = validate_rows({"a"}, [{"a": "1"}, {"a": ""}, {"a": "3"}])
    assert problems == ["Row 2: no value for 'a'."]


def test_validate_rows_ignores_columns_the_recipe_does_not_use() -> None:
    assert validate_rows({"a"}, [{"a": "1", "unused": ""}]) == []


def test_validate_rows_caps_a_long_problem_list() -> None:
    problems = validate_rows({"a"}, [{"a": ""}] * 50)
    assert len(problems) == 21
    assert problems[-1] == "...and 30 more problem(s)."


def test_validate_env_lists_only_missing_or_empty_variables() -> None:
    assert validate_env({"A", "B", "C"}, {"A": "x", "B": ""}) == [
        "The environment variable 'B' is not set.",
        "The environment variable 'C' is not set.",
    ]


# ---- column mapping ----


def test_normalize_key_ignores_case_spaces_and_punctuation() -> None:
    assert normalize_key("E-mail Address *") == "emailaddress"


def test_exact_names_match_first() -> None:
    assert suggest_mapping({"email"}, ["name", "email"]) == {"email": "email"}


@pytest.mark.parametrize("header", ["E-mail", "EMAIL", "e mail", "Email "])
def test_names_equal_after_normalising_are_matched(header: str) -> None:
    assert suggest_mapping({"email"}, [header]) == {"email": header}


def test_a_close_spelling_is_matched() -> None:
    assert suggest_mapping({"first_name"}, ["Firstnam", "age"]) == {"first_name": "Firstnam"}
    assert suggest_mapping({"telephone"}, ["Telephon"]) == {"telephone": "Telephon"}


def test_two_columns_that_normalise_the_same_are_not_guessed() -> None:
    assert suggest_mapping({"email"}, ["E-mail", "Email!"]) == {"email": None}


def test_no_match_is_none() -> None:
    assert suggest_mapping({"email"}, ["city", "age"]) == {"email": None}


def test_one_column_claimed_by_two_placeholders_is_left_unmatched() -> None:
    assert suggest_mapping({"mail", "email"}, ["e-mail"]) == {"mail": None, "email": None}


def test_no_headers_means_nothing_is_matched() -> None:
    assert suggest_mapping({"a"}, []) == {"a": None}


def test_apply_mapping_adds_placeholder_named_columns() -> None:
    rows = [{"E-mail": "a@x.io", "n": "Ann"}]
    assert apply_mapping(rows, {"email": "E-mail"}) == [
        {"E-mail": "a@x.io", "n": "Ann", "email": "a@x.io"}
    ]


def test_apply_mapping_rejects_an_unknown_column() -> None:
    with pytest.raises(DataError, match="'Nope'"):
        apply_mapping([{"a": "1"}], {"x": "Nope"})


def test_apply_mapping_on_no_rows_is_empty() -> None:
    assert apply_mapping([], {"x": "y"}) == []


# ---- secrets ----


@pytest.mark.parametrize(
    "name",
    ["password", "Password", "user_password", "PIN", "api_token", "aadhaar", "OTP", "db-passwd"],
)
def test_secret_looking_names_are_detected(name: str) -> None:
    assert is_secret_name(name) is True


@pytest.mark.parametrize("name", ["email", "name", "pinecone_city", "compass", "tokenizer"])
def test_ordinary_names_are_not_secrets(name: str) -> None:
    # "pinecone" and "compass" contain secret words only as substrings, not as whole words.
    assert is_secret_name(name) is False


def test_secret_values_collects_secret_cells_and_env_values() -> None:
    row = {"name": "Ann", "password": "hunter22", "pin": "12"}
    values = secret_values(row, ["TOKEN"], {"TOKEN": "abcdef"})
    assert values == ["hunter22", "abcdef"]  # longest first; "12" is too short to redact safely


def test_redact_replaces_every_secret() -> None:
    assert redact("login hunter22 then abcdef", ["abcdef", "hunter22"]) == "login *** then ***"
    assert redact("nothing here", ["hunter22"]) == "nothing here"
    assert redact("x", [""]) == "x"


def test_secret_columns_lists_the_secret_looking_headers() -> None:
    assert secret_columns(["name", "password", "pin"]) == ["password", "pin"]
