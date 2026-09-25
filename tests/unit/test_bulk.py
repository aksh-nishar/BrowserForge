"""Bulk runs: one recipe run per data row, against FakePageDriver."""

import json
from contextlib import nullcontext
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from browserforge.engine.bulk import BulkResult, RowResult, row_folder, row_label, run_bulk
from browserforge.engine.driver import DriverError
from browserforge.engine.recipe import Recipe, Step
from browserforge.engine.runner import DryRunError, run_recipe
from fixtures.fake_driver import FakePageDriver

NOW = datetime(2026, 9, 25, 10, 15, 30)


def step(step_type: str, name: str | None = None, **config: Any) -> Step:
    return Step(type=step_type, name=name, config=config)


def recipe(*steps: Step, **kwargs: Any) -> Recipe:
    return Recipe(url="https://x.io", steps=list(steps), **kwargs)


def fill_and_click() -> Recipe:
    return recipe(
        step("fill", "name", selector="#name", value="{{name}}"),
        step("click", "go", selector="{{button}}"),
    )


ROWS = [
    {"name": "Ann", "button": "#ok"},
    {"name": "Bob", "button": "#ok"},
    {"name": "Cy", "button": "#ok"},
]


def run(
    rec: Recipe,
    rows: list[dict[str, str]],
    fake: FakePageDriver,
    job_dir: Path,
    **kwargs: Any,
) -> BulkResult:
    return run_bulk(
        rec,
        rows,
        lambda: nullcontext(fake),
        job_dir,
        job_id="job1",
        recipe_name="demo",
        environ={},
        now=lambda: NOW,
        sleep=lambda _s: None,
        **kwargs,
    )


def sheet_rows(path: Path, sheet: str) -> list[dict[str, Any]]:
    worksheet = load_workbook(path)[sheet]
    table = list(worksheet.iter_rows(values_only=True))
    header = table[0]
    return [dict(zip((str(name) for name in header), row, strict=True)) for row in table[1:]]


# ---- all rows succeed ----


def test_every_row_runs_with_its_own_values(tmp_path: Path) -> None:
    fake = FakePageDriver()
    result = run(fill_and_click(), ROWS, fake, tmp_path)

    assert result.success is True
    assert [r.success for r in result.rows] == [True, True, True]
    assert fake.calls_to("fill") == [("#name", "Ann"), ("#name", "Bob"), ("#name", "Cy")]
    assert result.total == 3
    assert result.failed == 0


def test_each_row_gets_its_own_folder_with_a_run_log(tmp_path: Path) -> None:
    run(fill_and_click(), ROWS, FakePageDriver(), tmp_path)
    for number in (1, 2, 3):
        text = (tmp_path / row_folder(number) / "run_log.json").read_text("utf-8")
        log = json.loads(text)
        assert log["success"] is True
        assert log["job_id"] == "job1"


def test_a_summary_run_log_and_report_workbook_are_written(tmp_path: Path) -> None:
    result = run(fill_and_click(), ROWS, FakePageDriver(), tmp_path)

    summary = json.loads((tmp_path / "run_log.json").read_text("utf-8"))
    assert summary["total"] == 3
    assert [r["row"] for r in summary["rows"]] == [1, 2, 3]
    assert result.output_file == str(tmp_path / "output.xlsx")

    report = sheet_rows(tmp_path / "output.xlsx", "Run_Report")
    assert [r["Status"] for r in report] == ["OK", "OK", "OK"]
    assert [r["Row"] for r in report] == [1, 2, 3]
    assert report[0]["Details"] == "Ann / #ok"
    assert report[0]["Folder"] == "row_0001"


def test_the_start_page_is_opened_once_per_row(tmp_path: Path) -> None:
    fake = FakePageDriver()
    run(fill_and_click(), ROWS, fake, tmp_path)
    assert len(fake.calls_to("goto")) == 3


# ---- a failing row ----


def test_a_failed_row_is_recorded_and_the_rest_still_run(tmp_path: Path) -> None:
    rows = [
        {"name": "Ann", "button": "#ok"},
        {"name": "Bob", "button": "#bad"},
        {"name": "Cy", "button": "#ok"},
    ]
    fake = FakePageDriver()
    fake.set_missing("#bad")
    result = run(fill_and_click(), rows, fake, tmp_path)

    assert [r.success for r in result.rows] == [True, False, True]
    assert result.success is False
    assert result.failed == 1
    failed = result.rows[1]
    assert failed.failed_step == "go"
    assert failed.error is not None and "#bad" in failed.error
    assert [args[1] for args in fake.calls_to("fill")] == ["Ann", "Bob", "Cy"]


def test_a_failed_row_keeps_its_failure_screenshot_in_its_own_folder(tmp_path: Path) -> None:
    rows = [{"name": "Ann", "button": "#bad"}]
    fake = FakePageDriver()
    fake.set_missing("#bad")
    run(fill_and_click(), rows, fake, tmp_path)
    shots = list((tmp_path / "row_0001" / "screenshots").glob("failure_*.png"))
    assert len(shots) == 1


def test_the_report_lists_the_failure_with_its_reason(tmp_path: Path) -> None:
    rows = [{"name": "Ann", "button": "#bad"}, {"name": "Bob", "button": "#ok"}]
    fake = FakePageDriver()
    fake.set_missing("#bad")
    run(fill_and_click(), rows, fake, tmp_path)
    report = sheet_rows(tmp_path / "output.xlsx", "Run_Report")
    assert report[0]["Status"] == "Failed"
    assert report[0]["Failed step"] == "go"
    assert "#bad" in report[0]["Reason"]
    assert report[1]["Status"] == "OK"


def test_a_row_with_a_missing_value_fails_without_opening_the_browser(tmp_path: Path) -> None:
    rows = [{"name": "Ann"}, {"name": "Bob", "button": "#ok"}]  # first row lacks 'button'
    fake = FakePageDriver()
    result = run(fill_and_click(), rows, fake, tmp_path)
    assert result.rows[0].success is False
    assert "'button'" in (result.rows[0].error or "")
    assert result.rows[1].success is True
    assert len(fake.calls_to("goto")) == 1  # only the good row reached the browser


def test_a_crash_inside_one_row_does_not_stop_the_others(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    real = run_recipe
    calls = {"n": 0}

    def flaky(*args: Any, **kwargs: Any) -> Any:
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("boom")
        return real(*args, **kwargs)

    monkeypatch.setattr("browserforge.engine.bulk.run_recipe", flaky)
    result = run(fill_and_click(), ROWS, FakePageDriver(), tmp_path)
    assert [r.success for r in result.rows] == [True, False, True]
    assert "unexpected" in (result.rows[1].error or "")
    assert "boom" not in (result.rows[1].error or "")


def test_a_browser_that_cannot_start_aborts_the_whole_job(tmp_path: Path) -> None:
    def broken() -> Any:
        raise DriverError("no browser")

    with pytest.raises(DriverError):
        run_bulk(
            fill_and_click(),
            ROWS,
            broken,
            tmp_path,
            job_id="job1",
            environ={},
            now=lambda: NOW,
        )


# ---- scraped tables ----


def test_scraped_tables_from_every_row_are_merged_with_a_row_column(tmp_path: Path) -> None:
    fake = FakePageDriver()
    fake.set_table("#t", [["item"], ["pen"]])
    rec = recipe(
        step("fill", selector="#name", value="{{name}}"),
        step("extract_table", "Items", selector="#t"),
    )
    run(rec, [{"name": "Ann"}, {"name": "Bob"}], fake, tmp_path)
    assert sheet_rows(tmp_path / "output.xlsx", "Items") == [
        {"Row": 1, "item": "pen"},
        {"Row": 2, "item": "pen"},
    ]


# ---- secrets ----


def test_secret_cells_never_reach_logs_reports_or_row_folders(tmp_path: Path) -> None:
    rows = [{"name": "Ann", "password": "hunter22"}]
    rec = recipe(step("fill", "pw", selector="#pw", value="{{password}}"))
    fake = FakePageDriver()
    fake.fail_on("fill", DriverError("could not type hunter22"), times=None)
    result = run(rec, rows, fake, tmp_path)

    assert result.rows[0].success is False
    assert "hunter22" not in (result.rows[0].error or "")
    assert result.rows[0].label == "Ann"  # the password column is not shown
    for file in tmp_path.rglob("*.json"):
        assert "hunter22" not in file.read_text("utf-8")
    report = sheet_rows(tmp_path / "output.xlsx", "Run_Report")
    assert "hunter22" not in str(report)


def test_env_values_are_redacted_too(tmp_path: Path) -> None:
    rec = recipe(step("fill", "pw", selector="#pw", value="{{env:PORTAL_PW}}"))
    fake = FakePageDriver()
    fake.fail_on("fill", DriverError("could not type letmein99"), times=None)
    result = run_bulk(
        rec,
        [{"name": "Ann"}],
        lambda: nullcontext(fake),
        tmp_path,
        job_id="job1",
        environ={"PORTAL_PW": "letmein99"},
        now=lambda: NOW,
        sleep=lambda _s: None,
    )
    assert "letmein99" not in (result.rows[0].error or "")


# ---- dry run ----


def submit_recipe() -> Recipe:
    return recipe(
        step("fill", selector="#name", value="{{name}}"),
        step("click", "send", selector="#send", submit=True),
    )


def test_a_dry_run_never_clicks_submit_for_any_row(tmp_path: Path) -> None:
    fake = FakePageDriver()
    result = run(submit_recipe(), ROWS, fake, tmp_path, dry_run=True)
    assert result.success is True
    assert fake.clicks == []
    assert len(fake.screenshots) == 3


def test_a_dry_run_without_a_submit_step_is_refused_up_front(tmp_path: Path) -> None:
    fake = FakePageDriver()
    with pytest.raises(DryRunError):
        run(fill_and_click(), ROWS, fake, tmp_path, dry_run=True)
    assert fake.calls == []


# ---- progress callbacks and edge cases ----


def test_progress_callbacks_report_each_row(tmp_path: Path) -> None:
    events: list[tuple[str, int]] = []
    seen: list[RowResult] = []

    def on_end(row: RowResult) -> None:
        events.append(("end", row.row))
        seen.append(row)

    run(
        fill_and_click(),
        ROWS[:2],
        FakePageDriver(),
        tmp_path,
        on_row_start=lambda index, total, _label: events.append(("start", index * 10 + total)),
        on_row_end=on_end,
    )
    assert events == [("start", 12), ("end", 1), ("start", 22), ("end", 2)]
    assert all(r.success for r in seen)


def test_no_rows_is_not_a_success_but_still_writes_the_report(tmp_path: Path) -> None:
    result = run(fill_and_click(), [], FakePageDriver(), tmp_path)
    assert result.success is False
    assert result.rows == []
    assert (tmp_path / "output.xlsx").is_file()


def test_row_label_shows_a_few_plain_cells_and_hides_secrets() -> None:
    row = {"a": "one", "password": "pw123", "b": "", "c": "two", "d": "three", "e": "four"}
    assert row_label(row) == "one / two / three"


def test_row_label_is_shortened() -> None:
    label = row_label({"a": "x" * 200})
    assert len(label) == 60
    assert label.endswith("...")


def test_row_folder_names_sort_in_order() -> None:
    names = [row_folder(n) for n in (1, 2, 10, 1000)]
    assert names == sorted(names)
    assert names[0] == "row_0001"
