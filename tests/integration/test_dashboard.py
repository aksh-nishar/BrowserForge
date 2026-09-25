"""The real dashboard, in a real browser, against a live server that runs real jobs.

This automates the plan's manual test for Task 2.5: upload a recipe, click Run, watch the
status, see screenshots, download the Excel file. The job itself runs in the server's worker
thread with its own Chromium, exactly like production.
"""

import threading
import time
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pytest
import uvicorn
from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from browserforge.app import create_app
from browserforge.engine.driver import playwright_driver
from fixtures.local_site import PASSWORD, USERNAME

JOB_TIMEOUT_MS = 90_000


@dataclass
class LiveApp:
    url: str
    recipes: Path
    jobs: Path


@pytest.fixture
def live_app(tmp_path: Path) -> Iterator[LiveApp]:
    recipes, jobs = tmp_path / "recipes", tmp_path / "jobs"
    app = create_app(
        recipes_dir=recipes,
        jobs_dir=jobs,
        driver_factory=lambda: playwright_driver(timeout_ms=4_000),
    )
    # log_config=None: uvicorn's default config rewires the global "uvicorn.*" loggers for the
    # whole test process (propagate=False), which would leak into unrelated logging tests.
    config = uvicorn.Config(app, host="127.0.0.1", port=0, log_level="warning", log_config=None)
    server = uvicorn.Server(config)
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    deadline = time.monotonic() + 15
    while not server.started:
        assert time.monotonic() < deadline, "dashboard server did not start"
        time.sleep(0.05)
    port = server.servers[0].sockets[0].getsockname()[1]
    try:
        yield LiveApp(f"http://127.0.0.1:{port}", recipes, jobs)
    finally:
        server.should_exit = True
        thread.join(timeout=15)


def write_recipe(directory: Path, name: str, text: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{name}.yaml"
    path.write_text(text, encoding="utf-8")
    return path


def happy_recipe(site: str, password: str = PASSWORD) -> str:
    return f"""
url: {site}/login
login:
  username_selector: "#username"
  password_selector: "#password"
  submit_selector: "#submit"
  username: {USERNAME}
  password: "{password}"
  success_url_contains: /secure
  timeout_ms: 3000
steps:
  - type: screenshot
    name: logged_in
  - type: goto
    url: {site}/table?page=1
  - type: paginate
    name: Employees
    selector: "#employees"
    next_selector: "#next"
"""


def upload_and_select(page: Page, app: LiveApp, recipe: Path) -> None:
    page.goto(app.url)
    page.set_input_files("#recipe-file", str(recipe))
    page.click("#upload-btn")
    expect(page.locator("#message")).to_contain_text("Saved recipe")
    expect(page.locator("#recipe-select")).to_have_value(recipe.stem)


def test_full_flow_upload_run_watch_and_download(
    page: Page, live_app: LiveApp, site: str, tmp_path: Path
) -> None:
    recipe = write_recipe(tmp_path / "uploads", "employees", happy_recipe(site))
    upload_and_select(page, live_app, recipe)

    expect(page.locator("#run-btn")).to_be_enabled()
    page.click("#run-btn")

    # live status, then completion
    expect(page.locator("#status-card")).to_be_visible()
    expect(page.locator("#status-badge")).to_have_text("Complete", timeout=JOB_TIMEOUT_MS)
    expect(page.locator("#progress-text")).to_contain_text("Finished in")

    # the step log lists every step, all OK
    rows = page.locator("#steps-body tr")
    expect(rows).to_have_count(5)  # start page, login, screenshot, goto, paginate
    expect(page.locator("#steps-body td.result.bad")).to_have_count(0)
    expect(page.locator("#steps-body td.result.ok")).to_have_count(5)
    expect(rows.nth(1)).to_contain_text("login")
    expect(page.locator("#failure-box")).to_be_hidden()

    # results: screenshot gallery with images that really loaded
    expect(page.locator("#results-card")).to_be_visible()
    images = page.locator("#gallery img")
    expect(images).to_have_count(1)
    loaded = images.first.evaluate("img => img.complete && img.naturalWidth > 0")
    assert loaded is True

    # results: the Excel report downloads and contains all 6 scraped rows
    excel_link = page.locator("#downloads a", has_text="Excel report")
    expect(excel_link).to_have_count(1)
    with page.expect_download() as download_info:
        excel_link.click()
    saved = tmp_path / "report.xlsx"
    download_info.value.save_as(saved)
    sheet = load_workbook(saved)["Employees"]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[0] == ["ID", "Name", "Dept"]
    assert [row[1] for row in values[1:]] == ["Ann", "Bob", "Cy", "Di", "Ed", "Flo"]

    # the run log is offered too, and the recipe list shows the uploaded recipe
    expect(page.locator("#downloads a", has_text="Run log")).to_have_count(1)
    expect(page.locator("#recipe-select option")).to_have_text(["employees (3 steps)"])


def test_wrong_password_shows_a_failed_run_with_a_plain_reason(
    page: Page, live_app: LiveApp, site: str, tmp_path: Path
) -> None:
    recipe = write_recipe(tmp_path / "uploads", "bad_login", happy_recipe(site, "wrong"))
    upload_and_select(page, live_app, recipe)
    page.click("#run-btn")

    expect(page.locator("#status-badge")).to_have_text("Failed", timeout=JOB_TIMEOUT_MS)
    failure = page.locator("#failure-box")
    expect(failure).to_be_visible()
    expect(failure).to_contain_text("Login failed")
    expect(failure).to_contain_text("username and password")
    expect(page.locator("#steps-body td.result.bad")).to_have_count(1)
    expect(page.locator("#run-btn")).to_be_enabled()  # can run again

    # the automatic failure screenshot is in the gallery, labelled, and really loads
    expect(page.locator("#gallery figcaption.failure")).to_have_count(1)
    expect(page.locator("#gallery figcaption.failure")).to_contain_text("Failure screenshot")
    failure_image = page.locator("#gallery img").first
    assert failure_image.evaluate("img => img.complete && img.naturalWidth > 0") is True
    # and the message shown is the plain-language one, not the raw technical text
    assert "Timed out after" not in (failure.text_content() or "")


def test_invalid_recipe_upload_shows_the_validation_message(
    page: Page, live_app: LiveApp, tmp_path: Path
) -> None:
    bad = write_recipe(
        tmp_path / "uploads", "broken", "url: https://x.io\nsteps:\n  - type: clik\n"
    )
    page.goto(live_app.url)
    page.set_input_files("#recipe-file", str(bad))
    page.click("#upload-btn")
    message = page.locator("#message")
    expect(message).to_contain_text("unknown step type 'clik'")
    expect(message).to_have_class("error")
    expect(page.locator("#run-btn")).to_be_disabled()
    assert not (live_app.recipes / "broken.yaml").exists()


def test_upload_without_choosing_a_file_asks_for_one(page: Page, live_app: LiveApp) -> None:
    page.goto(live_app.url)
    page.click("#upload-btn")
    expect(page.locator("#message")).to_contain_text("Choose a .yaml recipe file first")


def test_page_starts_with_no_recipes_and_run_disabled(page: Page, live_app: LiveApp) -> None:
    page.goto(live_app.url)
    expect(page.locator("#recipe-select")).to_contain_text("No recipes yet")
    expect(page.locator("#run-btn")).to_be_disabled()
    expect(page.locator("#status-card")).to_be_hidden()
    expect(page.locator("#results-card")).to_be_hidden()


def test_untrusted_text_in_step_names_and_errors_is_never_run_as_html(
    page: Page, live_app: LiveApp, site: str, tmp_path: Path
) -> None:
    evil = "<img src=x onerror=window.__pwned=1>"
    recipe = write_recipe(
        tmp_path / "uploads",
        "xss",
        f"""
url: {site}/form
continue_on_error: true
steps:
  - type: wait_for_selector
    name: '{evil}'
    selector: "#does-not-exist"
    timeout_ms: 300
""",
    )
    upload_and_select(page, live_app, recipe)
    page.click("#run-btn")

    expect(page.locator("#status-badge")).to_have_text("Failed", timeout=JOB_TIMEOUT_MS)
    expect(page.locator("#steps-body")).to_contain_text(evil)  # shown as literal text
    expect(page.locator("#failure-box")).to_contain_text(evil)
    assert page.locator("#steps-body img, #failure-box img").count() == 0
    page.wait_for_timeout(300)  # give a bad onerror handler time to fire
    pwned: Any = page.evaluate("window.__pwned")
    assert pwned is None
