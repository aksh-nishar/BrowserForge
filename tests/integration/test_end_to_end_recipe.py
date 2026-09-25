"""The shipped example recipes, run start to finish on a real browser against the public
practice sites they were written for. Skipped automatically when those sites are unreachable."""

import json
import socket
from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook
from playwright.sync_api import Browser

from browserforge.engine.driver import PlaywrightDriver
from browserforge.engine.recipe import load_recipe
from browserforge.engine.runner import RunLog, run_recipe

RECIPES_DIR = Path(__file__).parents[2] / "recipes"
PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
PUBLIC_HOSTS = ("the-internet.herokuapp.com", "www.scrapethissite.com")


def _reachable(host: str) -> bool:
    try:
        socket.create_connection((host, 443), timeout=3).close()
    except OSError:
        return False
    return True


requires_public_sites = pytest.mark.skipif(
    not all(_reachable(host) for host in PUBLIC_HOSTS),
    reason="public practice sites are not reachable",
)


@pytest.fixture
def public_driver(browser: Browser) -> Iterator[PlaywrightDriver]:
    """Real sites are slower than the local test site, so a generous timeout."""
    context = browser.new_context(accept_downloads=True)
    driver = PlaywrightDriver(context.new_page(), timeout_ms=30_000)
    try:
        yield driver
    finally:
        driver.close()
        context.close()


def assert_all_steps_succeeded(log: RunLog) -> None:
    failures = [(s.name or s.type, s.human_error) for s in log.steps if not s.success]
    assert failures == []
    assert log.success is True


def screenshots_in(job_dir: Path) -> list[Path]:
    return sorted((job_dir / "screenshots").glob("*.png"))


# ---- Recipe A: login + scrape + paginate + Excel ----


@requires_public_sites
def test_recipe_a_login_scrape_paginate_and_export(
    public_driver: PlaywrightDriver, tmp_path: Path
) -> None:
    job_dir = tmp_path / "job-a"
    recipe = load_recipe(RECIPES_DIR / "example_login_scrape.yaml")
    log = run_recipe(recipe, public_driver, job_dir, recipe_name="example_login_scrape")

    assert_all_steps_succeeded(log)
    assert [s.type for s in log.steps] == [
        "goto",
        "login",
        "screenshot",
        "goto",
        "paginate",
        "screenshot",
    ]

    # output.xlsx: 3 pages x 25 teams = 75 data rows under a real header
    assert log.output_file == str(job_dir / "output.xlsx")
    sheet = load_workbook(job_dir / "output.xlsx")["Hockey_Teams"]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[0][0] == "Team Name"
    assert len(values) - 1 == 75
    assert all(row[0] for row in values[1:])  # every row has a team name
    assert len({f"{row[0]}|{row[1]}" for row in values[1:]}) == 75  # no page was read twice

    # the page cap cut the 24-page table short, and the log says so
    paginate = log.steps[4]
    assert paginate.data == {"row_count": 75, "pages": 3, "truncated": True}

    # at least N=2 real screenshots (after login, and of the table element)
    shots = screenshots_in(job_dir)
    assert len(shots) >= 2
    assert all(shot.read_bytes().startswith(PNG_SIGNATURE) for shot in shots)

    # run_log.json: valid, all steps succeeded, no credentials or table rows in it
    saved = json.loads((job_dir / "run_log.json").read_text(encoding="utf-8"))
    assert saved["success"] is True
    assert all(step["success"] for step in saved["steps"])
    text = json.dumps(saved)
    assert "SuperSecretPassword" not in text
    assert "Boston Bruins" not in text
