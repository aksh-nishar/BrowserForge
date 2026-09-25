"""Bulk runs, schedules and the `browserforge run` command against a real browser."""

import io
import json
import time
import urllib.parse
import urllib.request
from collections.abc import Callable
from contextlib import AbstractContextManager
from pathlib import Path
from typing import Any

from apscheduler.triggers.interval import IntervalTrigger
from openpyxl import load_workbook

from browserforge.cli import run_command
from browserforge.data_input import build_rows
from browserforge.engine.bulk import run_bulk
from browserforge.engine.driver import PlaywrightDriver, playwright_driver
from browserforge.engine.recipe import parse_recipe_text
from fixtures.live_server import LiveApp, register_recipe, running_app, submissions, write_recipe

PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
CSV = (
    "First Name,E-mail,plan\n"
    "Ann,ann@example.com,basic\n"
    "Bob,bob@example.com,pro\n"
    "Cy,cy@example.com,basic\n"
)
LOOK_RECIPE = "url: {site}/form\nsteps:\n  - type: screenshot\n    name: seen\n"


def real_browser() -> AbstractContextManager[PlaywrightDriver]:
    return playwright_driver(timeout_ms=3_000)


def report_rows(path: Path) -> list[dict[str, Any]]:
    table = list(load_workbook(path)["Run_Report"].iter_rows(values_only=True))
    header = [str(name) for name in table[0]]
    return [dict(zip(header, row, strict=True)) for row in table[1:]]


def http(app: LiveApp, method: str, path: str, form: dict[str, str] | None = None) -> Any:
    """One request to the live dashboard; returns the decoded JSON."""
    body = urllib.parse.urlencode(form).encode() if form is not None else None
    request = urllib.request.Request(app.url + path, method=method, data=body)
    with urllib.request.urlopen(request, timeout=15) as response:
        return json.loads(response.read())


def wait_until(check: Callable[[], Any], timeout_s: float = 60.0) -> Any:
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        found = check()
        if found:
            return found
        time.sleep(0.25)
    raise AssertionError("condition was not met in time")


def bulk_rows(site: str, sheet: str) -> tuple[Any, list[dict[str, str]]]:
    recipe = parse_recipe_text(register_recipe(site))
    return recipe, build_rows(recipe, raw=sheet.encode(), filename="people.csv", environ={})


# ---- bulk ----


def test_a_sheet_of_three_people_submits_three_real_forms(site: str, tmp_path: Path) -> None:
    recipe, rows = bulk_rows(site, CSV)
    before = len(submissions(site))

    result = run_bulk(
        recipe, rows, real_browser, tmp_path / "job", job_id="job1", recipe_name="reg", environ={}
    )

    assert result.success is True
    assert [r.success for r in result.rows] == [True, True, True]
    received = submissions(site)[before:]
    assert [(s["first_name"], s["email"], s["plan"]) for s in received] == [
        ("Ann", "ann@example.com", "basic"),
        ("Bob", "bob@example.com", "pro"),
        ("Cy", "cy@example.com", "basic"),
    ]
    report = report_rows(tmp_path / "job" / "output.xlsx")
    assert [r["Status"] for r in report] == ["OK", "OK", "OK"]
    assert report[0]["Details"] == "Ann / ann@example.com / basic"
    for number in (1, 2, 3):
        shots = list((tmp_path / "job" / f"row_{number:04d}" / "screenshots").glob("*.png"))
        assert shots
        assert shots[0].read_bytes().startswith(PNG_SIGNATURE)


def test_one_bad_row_fails_alone_with_a_screenshot_and_the_others_still_submit(
    site: str, tmp_path: Path
) -> None:
    sheet = (
        "first_name,email,plan\n"
        "Ann,ann@example.com,basic\n"
        "Bob,bob@example.com,gold\n"  # there is no 'gold' plan on the form
        "Cy,cy@example.com,pro\n"
    )
    recipe, rows = bulk_rows(site, sheet)
    before = len(submissions(site))

    result = run_bulk(
        recipe, rows, real_browser, tmp_path / "job", job_id="job2", recipe_name="reg", environ={}
    )

    assert [r.success for r in result.rows] == [True, False, True]
    assert result.success is False
    assert result.rows[1].failed_step == "plan"
    assert [s["first_name"] for s in submissions(site)[before:]] == ["Ann", "Cy"]
    failure_shots = list((tmp_path / "job" / "row_0002" / "screenshots").glob("failure_*.png"))
    assert len(failure_shots) == 1
    assert failure_shots[0].read_bytes().startswith(PNG_SIGNATURE)
    report = report_rows(tmp_path / "job" / "output.xlsx")
    assert [r["Status"] for r in report] == ["OK", "Failed", "OK"]
    assert report[1]["Failed step"] == "plan"


def test_a_test_run_of_a_sheet_submits_nothing(site: str, tmp_path: Path) -> None:
    recipe, rows = bulk_rows(site, CSV)
    before = len(submissions(site))
    result = run_bulk(
        recipe, rows, real_browser, tmp_path / "job", job_id="job3", environ={}, dry_run=True
    )
    assert result.success is True
    assert len(submissions(site)) == before


# ---- command line ----


def test_the_run_command_submits_one_form_per_row_and_exits_zero(site: str, tmp_path: Path) -> None:
    recipes = tmp_path / "recipes"
    write_recipe(recipes, "register", register_recipe(site))
    data = tmp_path / "people.csv"
    data.write_text(CSV, encoding="utf-8")
    out = io.StringIO()
    before = len(submissions(site))

    code = run_command(
        [
            "register",
            "--data",
            str(data),
            "--map",
            "first_name=First Name",
            "--map",
            "email=E-mail",
            "--recipes-dir",
            str(recipes),
            "--jobs-dir",
            str(tmp_path / "jobs"),
        ],
        out=out,
    )

    assert code == 0, out.getvalue()
    assert "Rows: 3 succeeded, 0 failed, of 3" in out.getvalue()
    assert [s["first_name"] for s in submissions(site)[before:]] == ["Ann", "Bob", "Cy"]
    (job_dir,) = (tmp_path / "jobs").iterdir()
    assert (job_dir / "output.xlsx").is_file()


def test_the_run_command_dry_run_submits_nothing(site: str, tmp_path: Path) -> None:
    recipes = tmp_path / "recipes"
    write_recipe(recipes, "register", register_recipe(site))
    data = tmp_path / "people.csv"
    data.write_text("first_name,email,plan\nAnn,ann@example.com,pro\n", encoding="utf-8")
    before = len(submissions(site))

    code = run_command(
        [
            "register",
            "--data",
            str(data),
            "--dry-run",
            "--recipes-dir",
            str(recipes),
            "--jobs-dir",
            str(tmp_path / "jobs"),
        ],
        out=io.StringIO(),
    )
    assert code == 0
    assert len(submissions(site)) == before


# ---- schedules ----


def test_a_schedule_fires_on_its_own_and_runs_a_real_job(site: str, tmp_path: Path) -> None:
    write_recipe(tmp_path / "recipes", "look", LOOK_RECIPE.format(site=site))
    # A 1-second trigger stands in for the cron time, so the test does not wait for the clock.
    with running_app(
        tmp_path, start_scheduler=True, trigger_factory=lambda _s: IntervalTrigger(seconds=1)
    ) as app:
        created = http(app, "POST", "/schedules", {"recipe_name": "look", "when": "hourly"})
        assert created["enabled"] is True

        def finished() -> dict[str, Any] | None:
            (entry,) = http(app, "GET", "/schedules")["schedules"]
            return entry if entry["last_status"] == "complete" else None

        entry = wait_until(finished)
        assert entry["last_error"] is None
        results = http(app, "GET", f"/jobs/{entry['last_job_id']}/results")
        assert [s["name"].endswith(".png") for s in results["screenshots"]] == [True]
        assert http(app, "GET", "/schedules")["running"] is True


def test_a_paused_schedule_stops_firing_and_can_be_deleted(site: str, tmp_path: Path) -> None:
    write_recipe(tmp_path / "recipes", "look", LOOK_RECIPE.format(site=site))
    with running_app(
        tmp_path, start_scheduler=True, trigger_factory=lambda _s: IntervalTrigger(seconds=1)
    ) as app:
        created = http(app, "POST", "/schedules", {"recipe_name": "look", "when": "hourly"})
        wait_until(lambda: http(app, "GET", "/schedules")["schedules"][0]["last_job_id"])

        http(app, "POST", f"/schedules/{created['id']}/pause")
        time.sleep(3)  # let a run that was already in flight finish
        settled = http(app, "GET", "/schedules")["schedules"][0]["last_job_id"]
        time.sleep(3)
        assert http(app, "GET", "/schedules")["schedules"][0]["last_job_id"] == settled

        assert http(app, "DELETE", f"/schedules/{created['id']}") == {"deleted": True}
        assert http(app, "GET", "/schedules")["schedules"] == []


def test_a_schedule_survives_a_server_restart(site: str, tmp_path: Path) -> None:
    write_recipe(tmp_path / "recipes", "look", LOOK_RECIPE.format(site=site))
    with running_app(tmp_path) as app:
        http(app, "POST", "/schedules", {"recipe_name": "look", "when": "daily 08:00"})
    with running_app(tmp_path, start_scheduler=True) as restarted:
        (entry,) = http(restarted, "GET", "/schedules")["schedules"]
        assert (entry["recipe_name"], entry["when"], entry["enabled"]) == (
            "look",
            "daily 08:00",
            True,
        )
        assert entry["next_run"] is not None
