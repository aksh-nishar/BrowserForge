"""Day 1 checkpoint: the whole engine (recipe -> runner -> actions -> Excel/log) on a real
browser. Runs against the local test site so it is deterministic and works offline."""

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from browserforge.engine.driver import PlaywrightDriver
from browserforge.engine.recipe import Recipe, Step, parse_recipe
from browserforge.engine.runner import run_recipe
from fixtures.local_site import PASSWORD, USERNAME

PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


def recipe_for(site: str, password: str = PASSWORD, **overrides: Any) -> Recipe:
    data: dict[str, Any] = {
        "url": f"{site}/login",
        "login": {
            "username_selector": "#username",
            "password_selector": "#password",
            "submit_selector": "#submit",
            "username": USERNAME,
            "password": password,
            "success_url_contains": "/secure",
            "timeout_ms": 2000,
        },
        "steps": [
            {"type": "screenshot", "name": "logged_in"},
            {"type": "goto", "url": f"{site}/table?page=1"},
            {
                "type": "paginate",
                "name": "Employees",
                "selector": "#employees",
                "next_selector": "#next",
            },
            {"type": "screenshot", "name": "table", "selector": "#employees"},
            {"type": "goto", "url": f"{site}/download"},
            {"type": "download", "selector": "#export", "save_as": "employees.csv"},
        ],
    }
    data.update(overrides)
    return parse_recipe(data)


def test_login_scrape_paginate_screenshot_download_and_excel(
    driver: PlaywrightDriver, site: str, tmp_path: Path
) -> None:
    job_dir = tmp_path / "job-1"
    log = run_recipe(recipe_for(site), driver, job_dir, recipe_name="day1")

    # every step succeeded
    failures = [(s.name or s.type, s.error) for s in log.steps if not s.success]
    assert failures == []
    assert log.success is True
    assert [s.type for s in log.steps] == [
        "goto",
        "login",
        "screenshot",
        "goto",
        "paginate",
        "screenshot",
        "goto",
        "download",
    ]

    # run_log.json: written, parseable, matches the returned log, holds no secrets or rows
    text = (job_dir / "run_log.json").read_text(encoding="utf-8")
    saved = json.loads(text)
    assert saved["success"] is True
    assert saved["job_id"] == "job-1"
    assert saved["recipe_name"] == "day1"
    assert PASSWORD not in text
    assert "Flo" not in text
    assert saved["steps"][4]["data"] == {"row_count": 6, "pages": 3, "truncated": False}

    # screenshots: real PNGs
    shots = sorted((job_dir / "screenshots").iterdir())
    assert len(shots) == 2
    assert all(shot.read_bytes().startswith(PNG_SIGNATURE) for shot in shots)

    # output.xlsx: all three pages of the table
    assert log.output_file == str(job_dir / "output.xlsx")
    sheet = load_workbook(job_dir / "output.xlsx")["Employees"]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values[0] == ["ID", "Name", "Dept"]
    assert [row[1] for row in values[1:]] == ["Ann", "Bob", "Cy", "Di", "Ed", "Flo"]

    # download: saved under the requested name
    assert (job_dir / "downloads" / "employees.csv").read_text() == "id,name\n1,Ann\n2,Bob\n"


def test_wrong_password_stops_the_run_with_a_clear_login_error(
    driver: PlaywrightDriver, site: str, tmp_path: Path
) -> None:
    job_dir = tmp_path / "job-2"
    log = run_recipe(recipe_for(site, password="wrong-password"), driver, job_dir)

    assert log.success is False
    assert log.stopped_early is True
    assert [s.type for s in log.steps] == ["goto", "login"]
    assert log.skipped_steps == 6
    login = log.steps[1]
    assert login.error is not None
    assert login.error.startswith("Login failed")
    assert "username and password" in login.error

    text = (job_dir / "run_log.json").read_text(encoding="utf-8")
    assert "wrong-password" not in text
    assert not (job_dir / "output.xlsx").exists()


def test_failed_step_with_continue_on_error_still_produces_the_rest(
    driver: PlaywrightDriver, site: str, tmp_path: Path
) -> None:
    recipe = recipe_for(site, continue_on_error=True)
    recipe.steps.insert(0, Step(type="click", config={"selector": "#nope"}))
    job_dir = tmp_path / "job-3"
    log = run_recipe(recipe, driver, job_dir)

    assert log.success is False  # one step failed
    assert log.stopped_early is False
    failed = [s for s in log.steps if not s.success]
    assert [s.type for s in failed] == ["click"]
    assert failed[0].error is not None
    assert "'#nope'" in failed[0].error
    assert (job_dir / "output.xlsx").exists()  # the later table step still ran
    assert (job_dir / "downloads" / "employees.csv").exists()
