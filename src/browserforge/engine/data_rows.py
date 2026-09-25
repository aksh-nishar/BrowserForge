"""Bulk data: spreadsheet rows -> values for `{{placeholders}}` in a recipe. No browser, no network.

A recipe can hold `{{name}}` placeholders (in `url`, `login` and step settings). Each row of a
CSV/XLSX sheet supplies the values for one run; the sheet's column headers are the names.
`{{env:NAME}}` reads an environment variable instead, which is how secrets (passwords, ID
numbers) get into a run without ever being written to a recipe, schedule or run file.
"""

import csv
import difflib
import io
import os
import re
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from browserforge.engine.recipe import Recipe, Step

MAX_ROWS = 1000
MAX_DATA_BYTES = 5_000_000
DATA_EXTENSIONS = (".csv", ".xlsx")
MAX_PROBLEMS_SHOWN = 20
MIN_SECRET_LENGTH = 3

_PLACEHOLDER = re.compile(r"\{\{\s*(env:)?([A-Za-z_][A-Za-z0-9_]*)\s*\}\}")
_SECRET_WORDS = frozenset(
    {
        "password",
        "passwd",
        "pass",
        "secret",
        "token",
        "otp",
        "cvv",
        "pin",
        "ssn",
        "aadhaar",
        "aadhar",
        "apikey",
        "credential",
        "credentials",
    }
)


class DataError(Exception):
    """The data file, a mapping, or a placeholder value is unusable. The message is for people."""


# ---------------------------------------------------------------------------
# Reading the sheet
# ---------------------------------------------------------------------------


def _cell_text(value: Any) -> str:
    """One spreadsheet cell as text: whole floats lose the '.0', dates become ISO dates."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _rows_from_table(table: Iterable[Sequence[Any]]) -> list[dict[str, str]]:
    """First row is the header. Blank rows and blank-header columns are dropped."""
    iterator = iter(table)
    header_row = next(iterator, None)
    if header_row is None:
        raise DataError("The data file is empty: it needs a header row with the column names.")

    headers = [_cell_text(cell) for cell in header_row]
    kept = [(index, name) for index, name in enumerate(headers) if name]
    names = [name for _, name in kept]
    if not names:
        raise DataError("The first row of the data file must hold the column names.")
    duplicates = sorted({name for name in names if names.count(name) > 1})
    if duplicates:
        raise DataError(f"The data file uses the column name(s) twice: {', '.join(duplicates)}.")

    rows: list[dict[str, str]] = []
    for raw in iterator:
        row = {name: _cell_text(raw[index]) if index < len(raw) else "" for index, name in kept}
        if not any(row.values()):
            continue  # a blank line in the sheet
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise DataError(f"The data file has more than {MAX_ROWS} rows. Split it into parts.")
    return rows


def parse_rows(raw: bytes, filename: str) -> list[dict[str, str]]:
    """Rows of a CSV or XLSX file's bytes, each a {column name: text} dict."""
    suffix = Path(filename).suffix.lower()
    if suffix not in DATA_EXTENSIONS:
        raise DataError("The data file must be a .csv or .xlsx file.")
    if len(raw) > MAX_DATA_BYTES:
        raise DataError(f"The data file is too large (limit {MAX_DATA_BYTES // 1_000_000} MB).")

    if suffix == ".csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise DataError(
                "The CSV file must be saved as UTF-8 text (in Excel: Save As > CSV UTF-8)."
            ) from None
        return _rows_from_table(csv.reader(io.StringIO(text)))

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:  # openpyxl raises many types (zip, xml, key errors) for bad files
        raise DataError("The .xlsx file could not be read. Is it a valid Excel file?") from None
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise DataError("The Excel file has no sheets.")
        return _rows_from_table(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def load_rows(path: str | Path) -> list[dict[str, str]]:
    """Rows of a CSV/XLSX file on disk."""
    file = Path(path)
    try:
        raw = file.read_bytes()
    except FileNotFoundError:
        raise DataError(f"Data file not found: {file}") from None
    except OSError as exc:
        raise DataError(f"Cannot read data file {file}: {exc}") from exc
    return parse_rows(raw, file.name)


# ---------------------------------------------------------------------------
# Finding and filling placeholders
# ---------------------------------------------------------------------------


def _map_strings(value: Any, fn: Callable[[str], str]) -> Any:
    if isinstance(value, str):
        return fn(value)
    if isinstance(value, dict):
        return {key: _map_strings(item, fn) for key, item in value.items()}
    if isinstance(value, list):
        return [_map_strings(item, fn) for item in value]
    return value


def _recipe_texts(recipe: Recipe) -> list[str]:
    found: list[str] = []

    def collect(text: str) -> str:
        found.append(text)
        return text

    _map_strings(recipe.url, collect)
    _map_strings(recipe.login, collect)
    for step in recipe.steps:
        _map_strings(step.config, collect)
    return found


def _scan(recipe: Recipe, env: bool) -> set[str]:
    names: set[str] = set()
    for text in _recipe_texts(recipe):
        for match in _PLACEHOLDER.finditer(text):
            if bool(match.group(1)) == env:
                names.add(match.group(2))
    return names


def find_placeholders(recipe: Recipe) -> set[str]:
    """Every `{{name}}` the recipe uses (column names it needs from the data)."""
    return _scan(recipe, env=False)


def find_env_vars(recipe: Recipe) -> set[str]:
    """Every `{{env:NAME}}` the recipe uses (environment variables it needs)."""
    return _scan(recipe, env=True)


def expand_recipe(
    recipe: Recipe,
    row: Mapping[str, str],
    environ: Mapping[str, str] | None = None,
) -> Recipe:
    """A copy of `recipe` with placeholders replaced by `row` values (and env variables).

    The original is untouched. A placeholder with no value raises DataError.
    """
    env = os.environ if environ is None else environ

    def fill(text: str) -> str:
        def replacement(match: re.Match[str]) -> str:
            name = match.group(2)
            if match.group(1):
                if name not in env:
                    raise DataError(f"The environment variable '{name}' is not set.")
                return env[name]
            if name not in row:
                raise DataError(f"There is no value for '{name}' in the data.")
            return row[name]

        return _PLACEHOLDER.sub(replacement, text)

    steps = [
        Step(
            type=step.type,
            name=step.name,
            config=_map_strings(step.config, fill),
            retries=step.retries,
            backoff_seconds=step.backoff_seconds,
        )
        for step in recipe.steps
    ]
    return replace(
        recipe,
        url=fill(recipe.url),
        login=_map_strings(recipe.login, fill),
        steps=steps,
    )


# ---------------------------------------------------------------------------
# Checking the data before any browser starts
# ---------------------------------------------------------------------------


def _capped(problems: list[str]) -> list[str]:
    if len(problems) <= MAX_PROBLEMS_SHOWN:
        return problems
    extra = len(problems) - MAX_PROBLEMS_SHOWN
    return [*problems[:MAX_PROBLEMS_SHOWN], f"...and {extra} more problem(s)."]


def validate_rows(placeholders: Iterable[str], rows: Sequence[Mapping[str, str]]) -> list[str]:
    """Plain-language problems that would break a run: no rows, missing columns, blank cells."""
    needed = sorted(placeholders)
    if not rows:
        return ["The data file has no rows below the header."]

    problems: list[str] = []
    present = [name for name in needed if name in rows[0]]
    problems.extend(
        f"The data has no column for '{name}'." for name in needed if name not in rows[0]
    )
    for number, row in enumerate(rows, start=1):
        problems.extend(
            f"Row {number}: no value for '{name}'." for name in present if not row.get(name, "")
        )
    return _capped(problems)


def validate_env(env_vars: Iterable[str], environ: Mapping[str, str] | None = None) -> list[str]:
    """Problems for environment variables the recipe needs but the machine does not have."""
    env = os.environ if environ is None else environ
    return [
        f"The environment variable '{name}' is not set."
        for name in sorted(env_vars)
        if not env.get(name)
    ]


# ---------------------------------------------------------------------------
# Matching sheet columns to placeholder names
# ---------------------------------------------------------------------------


def normalize_key(text: str) -> str:
    """Compare names ignoring case, spaces and punctuation ('E-mail Address' == 'email_address')."""
    return re.sub(r"[^a-z0-9]", "", text.lower())


def suggest_mapping(placeholders: Iterable[str], headers: Iterable[str]) -> dict[str, str | None]:
    """Guess which column feeds each placeholder. None means no safe match: the user decides.

    Exact names win; then names equal after ignoring case/punctuation; then a close spelling.
    Two candidates, or a column claimed by two placeholders, is left unmatched rather than guessed.
    """
    header_list = list(headers)
    by_key: dict[str, list[str]] = {}
    for header in header_list:
        by_key.setdefault(normalize_key(header), []).append(header)

    mapping: dict[str, str | None] = {}
    for placeholder in sorted(placeholders):
        if placeholder in header_list:
            mapping[placeholder] = placeholder
            continue
        key = normalize_key(placeholder)
        if key in by_key:
            candidates = by_key[key]
            mapping[placeholder] = candidates[0] if len(candidates) == 1 else None
            continue
        close = difflib.get_close_matches(key, list(by_key), n=2, cutoff=0.8)
        if len(close) == 1 and len(by_key[close[0]]) == 1:
            mapping[placeholder] = by_key[close[0]][0]
        else:
            mapping[placeholder] = None

    claimed: dict[str, int] = {}
    for column in mapping.values():
        if column is not None:
            claimed[column] = claimed.get(column, 0) + 1
    return {
        placeholder: (column if column is not None and claimed[column] == 1 else None)
        for placeholder, column in mapping.items()
    }


def apply_mapping(
    rows: Sequence[Mapping[str, str]], mapping: Mapping[str, str]
) -> list[dict[str, str]]:
    """Rows with each placeholder name filled from its mapped column (other columns kept)."""
    if rows:
        unknown = sorted({header for header in mapping.values() if header not in rows[0]})
        if unknown:
            raise DataError(f"The data has no column named {', '.join(repr(u) for u in unknown)}.")
    mapped: list[dict[str, str]] = []
    for row in rows:
        new = dict(row)
        for placeholder, header in mapping.items():
            new[placeholder] = row[header]
        mapped.append(new)
    return mapped


# ---------------------------------------------------------------------------
# Secrets
# ---------------------------------------------------------------------------


def is_secret_name(name: str) -> bool:
    """Does a column or placeholder name look like it holds a secret (password, token, PIN...)?"""
    lowered = name.lower()
    if "password" in lowered:
        return True
    return any(token in _SECRET_WORDS for token in re.split(r"[^a-z0-9]+", lowered))


def secret_values(
    row: Mapping[str, str],
    env_vars: Iterable[str] = (),
    environ: Mapping[str, str] | None = None,
) -> list[str]:
    """Values that must never appear in logs: secret-named cells plus used env variable values."""
    env = os.environ if environ is None else environ
    values = [value for name, value in row.items() if is_secret_name(name)]
    values.extend(env.get(name, "") for name in env_vars)
    return sorted({v for v in values if len(v) >= MIN_SECRET_LENGTH}, key=len, reverse=True)


def redact(text: str, secrets: Iterable[str]) -> str:
    """`text` with every secret value replaced by ***."""
    for secret in sorted(secrets, key=len, reverse=True):
        if secret:
            text = text.replace(secret, "***")
    return text


def secret_columns(headers: Iterable[str]) -> list[str]:
    """Column names that look like secrets (a saved data file must not carry them)."""
    return [name for name in headers if is_secret_name(name)]
