"""Excel export: list[dict] -> openpyxl Workbook, built in memory (no disk IO).

Scraped data is untrusted, so cell handling is defensive:
- strings are always stored as text, so a value like "=HYPERLINK(...)" never becomes a formula
- control characters Excel rejects are stripped
- values over Excel's 32,767-character cell limit are truncated
"""

from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_CELL_CHARS = 32_767
MAX_SHEET_NAME = 31
MIN_COLUMN_WIDTH = 8.0
MAX_COLUMN_WIDTH = 80.0
_INVALID_SHEET_CHARS = "[]:*?/\\"


def column_width(values: Iterable[object]) -> float:
    """Width for a column: longest line of any value + 2, clamped to [8, 80]."""
    longest = 0
    for value in values:
        text = "" if value is None else str(value)
        longest = max(longest, max((len(line) for line in text.splitlines()), default=0))
    return min(max(longest + 2.0, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)


def safe_sheet_name(name: str, taken: Iterable[str] = ()) -> str:
    """Make `name` a valid, unique Excel sheet name (Excel compares names case-insensitively)."""
    cleaned = "".join("_" if ch in _INVALID_SHEET_CHARS else ch for ch in name)
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", cleaned).strip().strip("'").strip()
    base = (cleaned or "Sheet")[:MAX_SHEET_NAME]

    used = {t.lower() for t in taken}
    candidate = base
    counter = 2
    while candidate.lower() in used:
        suffix = f"_{counter}"
        candidate = f"{base[: MAX_SHEET_NAME - len(suffix)]}{suffix}"
        counter += 1
    return candidate


def _cell_value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, (bool, int, float)):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    return text[:MAX_CELL_CHARS]


def _columns(rows: Sequence[Mapping[str, Any]]) -> list[str]:
    """Column names in order of first appearance across all rows."""
    seen: dict[str, None] = {}
    for row in rows:
        for key in row:
            seen.setdefault(key, None)
    return list(seen)


def _write_sheet(ws: Any, rows: Sequence[Mapping[str, Any]]) -> None:
    columns = _columns(rows)
    if not columns:
        return

    header = [_cell_value(name) for name in columns]
    table: list[list[Any]] = [header]
    table.extend([_cell_value(row.get(name)) for name in columns] for row in rows)

    for r_index, values in enumerate(table, start=1):
        for c_index, value in enumerate(values, start=1):
            cell = ws.cell(row=r_index, column=c_index, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # force text: never let scraped "=..." become a formula
            if r_index == 1:
                cell.font = Font(bold=True)

    for c_index in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(c_index)].width = column_width(
            row[c_index - 1] for row in table
        )
    ws.freeze_panes = "A2"


def build_workbook(sheets: Mapping[str, Sequence[Mapping[str, Any]]]) -> Workbook:
    """One sheet per key. First row is a bold header; columns are auto-sized.

    An empty mapping still yields a single empty sheet, since a workbook needs one.
    """
    workbook = Workbook()
    default = workbook.active
    assert default is not None
    workbook.remove(default)

    names: list[str] = []
    for requested, rows in sheets.items():
        name = safe_sheet_name(requested, taken=names)
        names.append(name)
        _write_sheet(workbook.create_sheet(title=name), rows)

    if not names:
        workbook.create_sheet(title="Sheet")
    return workbook


def save_workbook(workbook: Workbook, path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target
